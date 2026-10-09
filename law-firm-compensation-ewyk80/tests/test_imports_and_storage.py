"""Import wizard, duplicate detection, persistence, finalization and Excel export tests."""

from __future__ import annotations

import io
from decimal import Decimal

import openpyxl
import pandas as pd
import pytest

from compensation.database import ADMIN_PASSWORD_ENV, ReadOnlyYearError
from compensation.demo import load_demo
from compensation.exports import build_workbook
from compensation.imports import (find_duplicate_collections, import_dataframe, mapping_problems,
                                  normalize, read_upload, suggest_mapping)
from compensation.service import finalize_year, run, start_year_from_prior

CSV = b"""Payment Date,Client Name,Project ID,Invoice #,Payment ID,Payment Amount,Applied to Fees,Applied to Expenses
05/01/2025,Acme,M1,INV1,PMT-1,"$1,050.00",1000.00,50.00
05/02/2025,Acme,M1,INV2,PMT-2,500.00,500.00,0.00
05/01/2025,Acme,M1,INV1,PMT-1,"$1,050.00",1000.00,50.00
"""


def test_suggest_mapping_handles_timesolv_headings():
    raw = read_upload("p.csv", CSV)
    m = suggest_mapping("collections", list(raw.columns))
    assert m["collection_date"] == "Payment Date"
    assert m["matter_id"] == "Project ID"
    assert m["fee_amount"] == "Applied to Fees"
    assert m["expense_amount"] == "Applied to Expenses"
    assert m["invoice_id"] == "Invoice #"
    assert not mapping_problems("collections", m)


def test_validation_reports_bad_rows():
    bad = CSV + b"not a date,Acme,M1,INV3,PMT-3,10,abc,0\n"
    raw = read_upload("p.csv", bad)
    norm = normalize("collections", raw, suggest_mapping("collections", list(raw.columns)))
    problems = {(e["Row"], e["Field"]) for e in norm.errors}
    assert (5, "Collection date") in problems
    assert (5, "Amount allocated to professional fees (if exported)") in problems


# 7 ------------------------------------------------------------------------------
def test_duplicate_payments_are_detected(db):
    raw = read_upload("p.csv", CSV)
    mapping = suggest_mapping("collections", list(raw.columns))
    assert len(find_duplicate_collections(normalize("collections", raw, mapping).frame)) == 2
    res = import_dataframe(db, 2025, "collections", raw, mapping, "p.csv", CSV, "tester")
    assert res.imported == 3 and res.flagged_duplicates == 1
    stored = db.load_data("collections", 2025)
    assert int(stored["excluded"].sum()) == 1
    assert (stored["duplicate_of"] != "").sum() == 1
    # re-importing the same file never double counts
    again = import_dataframe(db, 2025, "collections", raw, mapping, "p.csv", CSV, "tester")
    assert again.imported == 0 and again.duplicates == 3
    assert len(db.load_data("collections", 2025)) == 3
    # the original file is preserved with its hash
    batches = db.batches(2025)
    assert batches["file_hash"].nunique() == 1
    assert db.batch_file(int(batches["batch_id"].iloc[0]))[1] == CSV


def test_saved_mapping_is_recalled(db):
    from compensation.imports import recall_or_suggest

    raw = read_upload("p.csv", CSV)
    custom = suggest_mapping("collections", list(raw.columns))
    custom["client"] = ""
    import_dataframe(db, 2025, "collections", raw, custom, "p.csv", CSV, "tester")
    recalled, origin = recall_or_suggest(db, "collections", list(raw.columns))
    assert origin == "saved" and recalled["client"] == ""


def test_demo_data_end_to_end(db):
    load_demo(db, 2025, "tester")
    inputs, res = run(db, 2025)
    assert len(inputs.partners) == 11
    assert res.reconciled
    assert sum(res.partner_summary["Total compensation"]) == Decimal("9500000.00")
    cats = set(res.exceptions["Category"])
    assert "Unassigned associate hours" in cats  # intentional unresolved exception
    assert "Duplicate payments" in cats
    assert res.metrics["fallback_collections"] > 0
    # split origination on M-1003
    od = res.origination_detail
    assert set(od[od["Matter ID"] == "M-1003"]["Partner"]) == {"Margaret Chen", "Robert Alvarez"}
    # Smith works for Okafor by default but for Petrova on M-1004
    att = res.attribution_detail
    smith = att[att["Timekeeper"] == "Daniel Smith"]
    assert set(smith[smith["Matter ID"] == "M-1004"]["Credited partner"]) == {"Elena Petrova"}
    assert set(smith[smith["Matter ID"] == "M-1001"]["Credited partner"]) == {"David Okafor"}
    jones = att[(att["Timekeeper"] == "Rachel Jones") & (att["Matter ID"] == "M-1007")]
    assert set(jones["Credited partner"]) == {"Michael Brennan", "Aisha Mohammed"}


# 14 -----------------------------------------------------------------------------
def test_excel_export_is_generated(db):
    load_demo(db, 2025, "tester")
    inputs, res = run(db, 2025)
    data = build_workbook(inputs, res, db.audit_log(2025))
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == [
        "Executive Summary", "Policy Inputs", "Partner Compensation", "EWYK Detail",
        "Origination Detail", "Working Credit Detail", "Associate-Matter Mappings",
        "TimeSolv Collections", "TimeSolv Time Entries", "Partner Expenses",
        "Expense Allocation", "Exceptions", "Reconciliation", "Audit Log"]
    pc = wb["Partner Compensation"]
    assert str(pc["G5"].value).startswith("=IF(")  # formula cells present
    assert pc.freeze_panes == "B5"
    assert "Pool" in wb.defined_names
    values = openpyxl.load_workbook(io.BytesIO(data), data_only=True)["Partner Compensation"]
    assert values.cell(4, 11).value == "Total compensation (application)"
    assert round(sum(values.cell(r, 11).value for r in range(5, 16)), 2) == 9500000.00


def test_finalized_year_is_read_only_and_reopen_requires_password(db, monkeypatch):
    load_demo(db, 2025, "tester")
    with pytest.raises(ValueError, match="Cannot finalize"):
        finalize_year(db, 2025, "tester")  # demo has an unresolved exception
    _, res = run(db, 2025)
    att = res.attribution_detail
    ids = att[att["Credited partner"] == "UNASSIGNED"]["Entry ID"].unique()
    db.save_table("exclusions", 2025, pd.DataFrame([{"entry_id": i, "reason": "test",
                                                     "entered_by": "t", "entered_at": ""}
                                                    for i in ids]), "tester")
    sid = finalize_year(db, 2025, "tester")
    snap = db.load_snapshot(sid)
    for key in ("policy", "imported_files", "column_mappings", "partner_roster",
                "matter_originators", "supervisory_mappings", "results", "audit_log",
                "created_at"):
        assert key in snap
    assert db.is_finalized(2025)
    with pytest.raises(ReadOnlyYearError):
        db.save_table("partners", 2025, pd.DataFrame(), "tester")
    monkeypatch.delenv(ADMIN_PASSWORD_ENV, raising=False)
    with pytest.raises(PermissionError):
        db.reopen_year(2025, "tester", "x", "fix")
    monkeypatch.setenv(ADMIN_PASSWORD_ENV, "s3cret")
    with pytest.raises(PermissionError):
        db.reopen_year(2025, "tester", "wrong", "fix")
    db.reopen_year(2025, "admin", "s3cret", "correct a mapping")
    assert not db.is_finalized(2025)
    # a new year starts from the prior year's setup
    start_year_from_prior(db, 2025, 2026, "tester")
    assert list(db.load_table("partners", 2026)["name"]) == list(db.load_table("partners", 2025)["name"])
    assert db.load_policy(2026).ewyk_pct == 80
